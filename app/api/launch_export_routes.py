"""Tenant-scoped CSV/XLSX exports for the active company's saved tenders."""
from __future__ import annotations

import asyncio
import csv
from io import BytesIO, StringIO
from urllib.parse import urlsplit, urlunsplit, parse_qsl

from fastapi import APIRouter, HTTPException, Query, Request, Response
from pydantic import ValidationError

from app.database import DatabaseError
from .organization_workflow_routes import (
    _account, _active_company, _shortlist_repository, _shortlist_data_error,
)
from .schemas import TenderDiscoveryItem

router = APIRouter(
    prefix="/api/v1/organizations/{organization_id}",
    tags=["launch-export"],
)

HEADERS = (
    "Source", "Tender ID", "Title", "Buyer", "Deadline (source timezone)",
    "Initial value", "Currency", "Region", "Preliminary fit (%)",
    "Score stage", "Original tender URL",
)
MAX_ROWS = 200
CSV_TYPE = "text/csv; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def safe_spreadsheet_text(value: object) -> str:
    """Avoid formula execution including whitespace/control-prefix bypass."""
    text = "" if value is None else str(value)
    if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def original_tender_url(value: str | None) -> str:
    """Conservative link rendering; never 'fix' malformed URLs or credentials."""
    try:
        url = urlsplit(value or "")
        if url.scheme.lower() != "https" or not url.hostname:
            return ""
        if url.username or url.password or url.port not in (None, 443):
            return ""
        if any(ord(c) < 32 or ord(c) == 127 for c in value):
            return ""
        if any(k.lower() in {"token", "api_key", "apikey", "access_token", "password", "secret"}
               for k, _ in parse_qsl(url.query, keep_blank_values=True)):
            return ""
        return urlunsplit(url)
    except (ValueError, TypeError):
        return ""


def export_rows(records) -> list[list[str]]:
    rows = []
    for record in records[:MAX_ROWS]:
        try:
            opportunity = TenderDiscoveryItem.model_validate_json(record.snapshot_json)
        except (ValueError, ValidationError):
            raise HTTPException(503, "Saved opportunity data is unavailable.") from None
        score = opportunity.preliminary_scoring
        rows.append([
            opportunity.source, opportunity.tender_number or opportunity.external_id,
            opportunity.title, opportunity.customer, opportunity.deadline,
            opportunity.initial_price, opportunity.currency, opportunity.region,
            score.fit_score if score else None, opportunity.analysis_stage,
            original_tender_url(opportunity.url),
        ])
    return [[safe_spreadsheet_text(value) for value in row] for row in rows]


def serialize_csv(rows: list[list[str]]) -> bytes:
    result = StringIO(newline="")
    writer = csv.writer(result, lineterminator="\r\n")
    writer.writerow(HEADERS)
    writer.writerows(rows)
    return ("\ufeff" + result.getvalue()).encode("utf-8")


def serialize_xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Saved tenders"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="19304F")
    for n, width in enumerate((17, 23, 52, 35, 28, 19, 13, 24, 18, 20, 66), 1):
        sheet.column_dimensions[get_column_letter(n)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"  # Never interpret an untrusted tender as a formula.
        url = row[10].value
        if url and url.startswith("https://"):
            row[10].hyperlink = url
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@router.get("/shortlist/export")
async def export_shortlist(
    organization_id: int, request: Request,
    format: str = Query(default="csv", pattern="^(csv|xlsx)$"),
):
    account = await _account(request)
    company = await _active_company(request, account, organization_id)
    if company is None:
        raise HTTPException(409, "Create and activate a company first.")
    repo = _shortlist_repository(request)
    try:
        records = await asyncio.to_thread(
            repo.list_saved_opportunities_for_organization,
            account.id, organization_id, company.id, MAX_ROWS,
        )
    except DatabaseError as error:
        _shortlist_data_error(error)
    rows = export_rows(records)
    extension = format
    payload = serialize_csv(rows) if format == "csv" else serialize_xlsx(rows)
    return Response(
        content=payload,
        media_type=CSV_TYPE if format == "csv" else XLSX_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="valyqon-saved.{extension}"',
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
