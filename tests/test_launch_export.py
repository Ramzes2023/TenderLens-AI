"""CSV/XLSX security and tenant boundaries for launch exports."""
import asyncio
import csv
from io import BytesIO, StringIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

import app.api.launch_export_routes as exports


class ScopedRepository:
    def __init__(self):
        self.calls = []
        self.audit_calls = []

    def record_shortlist_export_for_organization(self, **kwargs):
        self.audit_calls.append(kwargs)

    def list_saved_opportunities_for_organization(self, *args):
        self.calls.append(args)
        return [SimpleNamespace(snapshot_json=(
            '{"source":"eis","external_id":"001","title":"=2+2",'
            '"url":"https://zakupki.gov.ru/notice/001",'
            '"reasons":[],"metadata_analysis":{},"preliminary_scoring":{"profile_name":"Test company","profile_version":"1","fit_score":82.0,"scorable_weight":100,"completeness_percent":80,"criteria":[]}}'
        ))]


def test_csv_quotes_formulas_and_keeps_external_ids():
    repo = ScopedRepository()
    rows = exports.export_rows(repo.list_saved_opportunities_for_organization(1, 2, 3, 200))
    assert rows[0][1] == "001"
    assert rows[0][2] == "'=2+2"
    payload = exports.serialize_csv(rows)
    parsed = list(csv.reader(StringIO(payload.decode("utf-8-sig"))))
    assert parsed[0] == list(exports.HEADERS)
    assert parsed[1][2] == "'=2+2"
    assert parsed[1][-1] == "https://zakupki.gov.ru/notice/001"


@pytest.mark.parametrize("raw", [
    "http://example.com", "javascript:alert(1)", "https://user:pass@example.com/",
    "https://example.com:8443/path", "https://example.com/?api_key=secret",
    "https://example.com/\nmalicious",
])
def test_original_url_fail_closed(raw):
    assert exports.original_tender_url(raw) == ""


def test_spreadsheet_formula_and_safe_text():
    assert exports.safe_spreadsheet_text("  \t@SUM(A1)") == "'  \t@SUM(A1)"
    assert exports.safe_spreadsheet_text("-2") == "'-2"
    assert exports.safe_spreadsheet_text("Normal buyer") == "Normal buyer"
    data = exports.serialize_xlsx([["'=SUM(1)", *[""] * 10]])
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(data), read_only=False).active
    assert ws["A2"].data_type != "f"


def test_api_uses_account_organization_and_active_company(monkeypatch):
    repo = ScopedRepository()
    async def account(_request):
        return SimpleNamespace(id=52)
    async def company(_request, _account, _org):
        return SimpleNamespace(id=77)
    monkeypatch.setattr(exports, "_account", account)
    monkeypatch.setattr(exports, "_active_company", company)
    monkeypatch.setattr(exports, "_shortlist_repository", lambda _: repo)
    response = asyncio.run(exports.export_shortlist(21, SimpleNamespace(), "csv"))
    assert repo.calls == [(52, 21, 77, 200)]
    assert repo.audit_calls == [dict(account_id=52, organization_id=21,
                                     company_id=77, format='csv', item_count=1)]
    assert response.headers["Cache-Control"] == "private, no-store"
    assert response.headers["Content-Disposition"].endswith('valyqon-saved.csv"')


def test_api_requires_active_company(monkeypatch):
    async def account(_request):
        return SimpleNamespace(id=52)
    async def no_company(_request, _account, _org):
        return None
    monkeypatch.setattr(exports, "_account", account)
    monkeypatch.setattr(exports, "_active_company", no_company)
    with pytest.raises(HTTPException) as error:
        asyncio.run(exports.export_shortlist(21, SimpleNamespace(), "csv"))
    assert error.value.status_code == 409
